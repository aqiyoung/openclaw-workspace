from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

BG = RGBColor(0x1a, 0x1a, 0x2e)
ACCENT = RGBColor(0x00, 0xd2, 0xff)
WHITE = RGBColor(0xff, 0xff, 0xff)
GRAY = RGBColor(0xcc, 0xcc, 0xcc)
ORANGE = RGBColor(0xff, 0x6b, 0x35)

def bg(slide):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = BG

def txt(slide, l, t, w, h, text, sz=18, c=WHITE, b=False, a=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    p = tf.paragraphs[0]; p.text = text
    p.font.size = Pt(sz); p.font.color.rgb = c; p.font.bold = b; p.alignment = a
    return tf

def bullets(slide, l, t, w, h, items, sz=16, c=WHITE):
    box = slide.shapes.add_textbox(Inches(l), Inches(t), Inches(w), Inches(h))
    tf = box.text_frame; tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item; p.font.size = Pt(sz); p.font.color.rgb = c; p.space_after = Pt(8)

# Slide 1: Title
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 1, 1.5, 11, 1.5, "OpenClaw", 54, ACCENT, True, PP_ALIGN.CENTER)
txt(s, 1, 3.2, 11, 1, "你的私人 AI 助手平台", 32, WHITE, False, PP_ALIGN.CENTER)
txt(s, 1, 4.5, 11, 1, "多通道接入 · 多模型支持 · 技能无限扩展", 20, GRAY, False, PP_ALIGN.CENTER)
txt(s, 1, 6.2, 11, 0.5, "docs.openclaw.ai | github.com/openclaw/openclaw", 14, GRAY, False, PP_ALIGN.CENTER)

# Slide 2: What
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 0.8, 0.5, 11, 1, "什么是 OpenClaw？", 36, ACCENT, True)
bullets(s, 0.8, 1.8, 11, 5, [
    "一个开源的 AI 助手平台，部署在你自己的服务器上",
    "支持多通道接入：微信、飞书、Telegram、Discord、WhatsApp、Slack...",
    "支持多模型：OpenRouter、OpenAI、Claude、DeepSeek、龙猫、腾讯元宝...",
    "技能系统（Skills）：通过安装技能无限扩展能力",
    "数据私有：所有对话、配置、记忆都存储在你的设备上",
    "开源免费：MIT 协议，完全自由使用和修改",
], 20)

# Slide 3: Channels
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 0.8, 0.5, 11, 1, "全平台接入", 36, ACCENT, True)
bullets(s, 0.8, 1.8, 5.5, 5, [
    "微信 — 扫码登录，直接对话",
    "Telegram — Bot 接入",
    "飞书 — 企业级 IM 集成",
    "Discord — 社区友好",
    "WhatsApp — 国际化通讯",
    "Slack — 团队协作",
], 20)
bullets(s, 7, 1.8, 5.5, 5, [
    "iMessage — Apple 生态",
    "Microsoft Teams",
    "Google Chat",
    "IRC — 经典协议",
    "LINE / Signal / Matrix",
    "更多通道持续添加中...",
], 20)
txt(s, 0.8, 5.8, 11, 1, "一个 AI 助手，所有平台同步响应，统一记忆", 22, ORANGE, True, PP_ALIGN.CENTER)

# Slide 4: Models
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 0.8, 0.5, 11, 1, "多模型支持", 36, ACCENT, True)
bullets(s, 0.8, 1.8, 11, 5.5, [
    "OpenRouter — 自动选择最优模型",
    "OpenAI (GPT-4o / o1 / o3)",
    "Anthropic (Claude 3.5 / 4)",
    "DeepSeek — 国产推理模型",
    "美团龙猫 (LongCat) — 免费/低价",
    "腾讯元宝 (混元) — 国产大模型",
    "Xiaomi MiMo — 小米模型",
    "Moonshot / Kimi / 百度千帆 / 火山引擎...",
    "Ollama — 本地部署，完全离线",
], 18)
txt(s, 0.8, 6.2, 11, 0.5, "可随时切换模型，支持 per-session 模型覆盖", 20, ORANGE, True, PP_ALIGN.CENTER)

# Slide 5: Skills
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 0.8, 0.5, 11, 1, "技能系统 — 无限扩展", 36, ACCENT, True)
bullets(s, 0.8, 1.8, 5.5, 5, [
    "飞书文档 — 创建/读取/编辑",
    "飞书日历 — 日程管理",
    "飞书任务 — 任务清单",
    "多维表格 — 数据管理",
    "网页搜索 — 实时信息",
    "天气查询",
], 20)
bullets(s, 7, 1.8, 5.5, 5, [
    "IMA 知识库 — 笔记/文档管理",
    "公众号写作 — 全流程创作",
    "Word/Excel — 文档处理",
    "代码执行 — Python/Shell",
    "网页抓取 — 内容提取",
    "更多技能：clawhub.com",
], 20)
txt(s, 0.8, 5.8, 11, 1, "一键安装，即装即用，技能库持续增长", 22, ORANGE, True, PP_ALIGN.CENTER)

# Slide 6: Memory
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 0.8, 0.5, 11, 1, "记忆系统 — 越用越懂你", 36, ACCENT, True)
bullets(s, 0.8, 1.8, 11, 4, [
    "短期记忆：当前会话上下文",
    "长期记忆：MEMORY.md 持久化存储",
    "每日记录：memory/YYYY-MM-DD.md 自动归档",
    "语义搜索：跨记忆文件智能检索",
    "偏好学习：记住你的习惯和喜好",
    "项目追踪：记住你正在做的事情",
], 20)
txt(s, 0.8, 5.5, 11, 1, "每次对话都是在积累，AI 越来越了解你", 22, ORANGE, True, PP_ALIGN.CENTER)

# Slide 7: Self-hosted
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 0.8, 0.5, 11, 1, "自主部署 — 数据完全掌控", 36, ACCENT, True)
bullets(s, 0.8, 1.8, 11, 4, [
    "部署在自己的服务器/NAS/树莓派上",
    "支持 systemd 服务管理，开机自启",
    "Gateway 本地运行，也可暴露到公网",
    "Tailscale / Nginx 反代 / HTTPS",
    "所有数据本地存储，不经过第三方",
    "支持备份和恢复",
], 20)
txt(s, 0.8, 5.5, 11, 1, "你的 AI，你的数据，你的规则", 22, ORANGE, True, PP_ALIGN.CENTER)

# Slide 8: Summary
s = prs.slides.add_slide(prs.slide_layouts[6]); bg(s)
txt(s, 1, 1, 11, 1, "OpenClaw 能做什么？", 40, ACCENT, True, PP_ALIGN.CENTER)
bullets(s, 1.5, 2.5, 10, 4.5, [
    "在微信/飞书/Telegram 上同时和你对话",
    "帮你写公众号文章、处理 Word/Excel 文档",
    "管理飞书日历、任务、文档、知识库",
    "搜索网页、抓取内容、分析数据",
    "执行代码、自动化运维任务",
    "记住你的偏好，越用越聪明",
    "一切都运行在你自己的服务器上",
], 22)
txt(s, 1, 6.3, 11, 0.8, "开源 · 免费 · 私有 · 可扩展", 28, ORANGE, True, PP_ALIGN.CENTER)

out = "/home/liyang/.openclaw/workspace/OpenClaw介绍.pptx"
prs.save(out)
print(f"Saved: {out}")
